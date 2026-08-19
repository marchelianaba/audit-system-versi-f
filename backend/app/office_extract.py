"""Ekstraksi teks dari dokumen Office (Word .docx, Excel .xlsx) → list halaman.

Dipakai lapisan digest agar upload Word/Excel diterima (bukan hanya PDF).
Deterministik, tanpa LLM. Aman dipanggil dari thread (murni Python).
"""
from __future__ import annotations

from pathlib import Path

# Suffix yang ditangani modul ini.
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
OFFICE_SUFFIXES = DOCX_SUFFIXES | XLSX_SUFFIXES


def extract_docx_pages(path: str | Path) -> list[str]:
    """Teks dari .docx: paragraf + isi tabel (baris digabung ' | '). Satu blob."""
    try:
        import docx  # python-docx
    except ImportError:
        return []
    try:
        d = docx.Document(str(path))
    except Exception:  # noqa: BLE001 — korup/terenkripsi
        return []
    parts: list[str] = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in getattr(d, "tables", []):
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [text] if text else []


def extract_xlsx_pages(path: str | Path) -> list[str]:
    """Teks dari .xlsx: satu 'halaman' per sheet, baris = sel non-kosong ' | '."""
    try:
        import openpyxl
    except ImportError:
        return []
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return []
    pages: list[str] = []
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip() != ""]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                pages.append(f"[Sheet: {ws.title}]\n" + "\n".join(rows))
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return pages


def extract_office_pages(path: str | Path) -> list[str] | None:
    """Router: return list halaman bila .docx/.xlsx; None bila bukan Office
    (pemanggil lanjut ke ekstraktor lain, mis. PDF)."""
    suf = Path(path).suffix.lower()
    if suf in DOCX_SUFFIXES:
        return extract_docx_pages(path)
    if suf in XLSX_SUFFIXES:
        return extract_xlsx_pages(path)
    return None
